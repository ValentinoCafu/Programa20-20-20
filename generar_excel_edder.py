import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def crear_excel_requerimiento_edder():
    wb = Workbook()
    
    # Colores corporativos Tecsur
    NAVY_HEADER = '002D62'
    GOLD_ACCENT = 'F59E0B'
    GRAY_LIGHT = 'F1F5F9'
    GRAY_BORDER = 'CBD5E1'
    TEXT_DARK = '0F172A'
    
    font_title = Font(name='Segoe UI', size=14, bold=True, color='FFFFFF')
    font_subtitle = Font(name='Segoe UI', size=10, italic=True, color='E2E8F0')
    font_header = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
    font_data = Font(name='Segoe UI', size=9, color=TEXT_DARK)
    font_data_bold = Font(name='Segoe UI', size=9, bold=True, color=TEXT_DARK)
    font_note = Font(name='Segoe UI', size=9, italic=True, color='475569')
    
    fill_navy = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type='solid')
    fill_subnavy = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    fill_gold = PatternFill(start_color=GOLD_ACCENT, end_color=GOLD_ACCENT, fill_type='solid')
    fill_highlight = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid') # Amarillo claro para campos a completar
    
    thin_border = Border(
        left=Side(style='thin', color=GRAY_BORDER),
        right=Side(style='thin', color=GRAY_BORDER),
        top=Side(style='thin', color=GRAY_BORDER),
        bottom=Side(style='thin', color=GRAY_BORDER)
    )
    
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_wrap_left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    supervisores = [
        ('sup-01', 'CARDENAS YNGA MARTHA ELENA', 'SUPERVISOR C', '4070', 'PLANIFICACION OBRAS Y MANTTO BT'),
        ('sup-02', 'MELGAREJO DIAZ STEVEN ALFREDO', 'SUPERVISOR C', '4070', 'PLANIFICACION OBRAS Y MANTTO BT'),
        ('sup-03', 'YATACO NAVARRO KEWIN HAROLD STEVE', 'SUPERVISOR C', '4070', 'PLANIFICACION OBRAS Y MANTTO BT'),
        ('sup-04', 'AYMA CHAUPIS EDSON LEONARDO', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-05', 'YAPU ABARCA LARRY ERICK', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-06', 'RAMOS AUQUE LUIS ALBERTO', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-07', 'DIBURCIO RAMOS MARCO ANTONIO', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-08', 'HERMITANO LEON SMITH FRANCKLIN', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-09', 'QUISPE ALFARO BLADIMIRO', 'SUPERVISOR B', '4301', 'OBRAS BAJA TENSION'),
        ('sup-10', 'VASCO MALLMA DENIS MARXS', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-11', 'HERNANDEZ LOPEZ FREDDY ANTONIO', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-12', 'VALLEJOS HOLGUIN CESAR ELOY', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-13', 'ESPIRILLA CACERES MARCO ANTONIO', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-14', 'VALENCIA AZULA SAMUEL ALEJANDRO', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-15', 'ELIAS MANRIQUE ALEXANDER ALEXIS', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-16', 'DOMINGUEZ RAMIREZ HENRY', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-17', 'CHIRITO MILLONES VICTOR JUAN RODOLFO', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-18', 'REANO LEON HANS ERNESTO', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-19', 'ARANGO PENA RAUL ROLANDO', 'SUPERVISOR JUNIOR', '4301', 'OBRAS BAJA TENSION'),
        ('sup-20', 'VILLA DIAZ JOEL ALFREDO', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-21', 'CHILENO DIAZ ELVIS JOSEPH', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-22', 'QUIMPER REYNOSO IVAN', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-23', 'RAMIREZ QUISPE LUIS ANTONNY', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION'),
        ('sup-24', 'CARBAJAL VILLASANA JOEL', 'SUPERVISOR C', '4301', 'OBRAS BAJA TENSION')
    ]

    contratistas = [
        'ALIHCOM BUILD SAC', 'ALVIMAR', 'CAMPOS LATAM', 'COSORCIO CARTER RENSA',
        'ENCOSSA SAC', 'EOS', 'GRUCOIN S.A.C.', 'GRUPO ACUARIUS',
        'GRUPO ENERGIA', 'KIMSA', 'NATURAL GAS COMPANY', 'PRODUCSA'
    ]

    # ==========================================
    # HOJA 0: Presentación y Alcance
    # ==========================================
    ws0 = wb.active
    ws0.title = "0. Instrucciones & Alcance"
    ws0.views.sheetView[0].showGridLines = True
    
    ws0.merge_cells("A1:G1")
    ws0["A1"] = "PROGRAMA DE EVALUACIÓN OPERATIVA Y CONDUCTUAL 20-20-20 (TECSUR / LUZ DEL SUR)"
    ws0["A1"].font = font_title
    ws0["A1"].fill = fill_navy
    ws0["A1"].alignment = align_center
    ws0.row_dimensions[1].height = 35

    ws0.merge_cells("A2:G2")
    ws0["A2"] = "REQUERIMIENTO OFICIAL DE MATRIZ DE ASIGNACIONES — DEPARTAMENTO OBRAS DE BAJA TENSIÓN"
    ws0["A2"].font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    ws0["A2"].fill = fill_subnavy
    ws0["A2"].alignment = align_center
    ws0.row_dimensions[2].height = 24

    instrucciones_data = [
        ("DESTINATARIO:", "Ing. Edder — Jefe de Departamento Obras de Baja Tensión (Tecsur S.A.)"),
        ("PROPÓSITO:", "Configurar la matriz de relacionamiento de los 24 supervisores para garantizar evaluaciones 100% objetivas, trazables y sin sesgos."),
        ("PRINCIPIO CLAVE:", "Ningún evaluador calificará a un supervisor con el que no trabaje directamente en campo o en gestión de proyectos."),
        ("LOGIN DE CAMPO:", "Los Capataces y Coordinadores Contratistas ingresarán a la plataforma web mediante su DNI (8 dígitos). Es mandatorio registrar el DNI fidedigno de cada uno."),
        ("ESTRUCTURA DE HOJAS:", "Este libro contiene 4 hojas de relevamiento más el directorio oficial de contratistas:"),
        ("  • Hoja 1 (Jefes de Sección):", "Asignar a cuál de los 3 Jefes de Sección reporta cada uno de los 24 supervisores de Obras BT y Planificación."),
        ("  • Hoja 2 (Coordinadores Contratista):", "Vincular qué coordinadores de las 12 contratistas interactúan con cada supervisor (DNI y teléfono)."),
        ("  • Hoja 3 (Capataces de Campo):", "Mapear los capataces de cuadrilla asignados habitualmente a cada supervisor (DNI y tipo de cuadrilla)."),
        ("  • Hoja 4 (Luz del Sur):", "Registrar los Inspectores y Proyectistas de Luz del Sur asignados por zona/sector operativo a cada supervisor."),
        ("  • Hoja 5 (Contratistas Homologadas):", "Directorio oficial de las 12 empresas contratistas homologadas por Tecsur y LDS."),
        ("PLAZO SUGERIDO:", "Completar los campos sombreados en amarillo claro y devolver este archivo para su carga al sistema web."),
        ("APLICATIVO EN VIVO:", "https://valentinocafu.github.io/growthbro/ (Plataforma oficial activa multidispositivo)")
    ]

    row_idx = 4
    for label, desc in instrucciones_data:
        ws0.cell(row=row_idx, column=1, value=label).font = font_data_bold
        ws0.cell(row=row_idx, column=1).alignment = align_left
        ws0.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=7)
        c = ws0.cell(row=row_idx, column=2, value=desc)
        c.font = font_data
        c.alignment = align_wrap_left
        ws0.row_dimensions[row_idx].height = 24
        row_idx += 1

    # ==========================================
    # HOJA 1: Supervisores vs Jefes de Sección
    # ==========================================
    ws1 = wb.create_sheet(title="1. Sup_vs_Jefes_Seccion")
    ws1.views.sheetView[0].showGridLines = True
    
    headers_1 = [
        "N°", "ID SUPERVISOR", "APELLIDOS Y NOMBRES", "CARGO ESCALERA", 
        "CR", "UNIDAD / ESPECIALIDAD", "JEFE DE SECCIÓN ASIGNADO (COMPLETAR)", 
        "ZONA / FRENTE OPERATIVO", "ESTADO"
    ]
    
    ws1.append(headers_1)
    ws1.row_dimensions[1].height = 28
    for col_num in range(1, len(headers_1) + 1):
        cell = ws1.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = align_center

    for idx, (sid, nombre, cargo, cr, crdesc) in enumerate(supervisores, start=1):
        # Sugerencia inicial lógica según CR
        sugerencia_jefe = "Jefe Sección 3 (Planificación BT)" if cr == '4070' else ("Jefe Sección 1" if idx % 2 == 0 else "Jefe Sección 2")
        zona_sug = "Sede San Juan / Planificación" if cr == '4070' else "Lima Sur / Obras BT"
        
        row_vals = [
            idx, sid, nombre, cargo, cr, crdesc, sugerencia_jefe, zona_sug, "Activo en Campo"
        ]
        ws1.append(row_vals)
        r = idx + 1
        ws1.row_dimensions[r].height = 20
        is_even = (idx % 2 == 0)
        
        for c in range(1, len(headers_1) + 1):
            cell = ws1.cell(row=r, column=c)
            cell.font = font_data
            cell.border = thin_border
            if c in (1, 2, 5):
                cell.alignment = align_center
            else:
                cell.alignment = align_left
            
            # Resaltar la columna editable por Edder
            if c == 7:
                cell.fill = fill_highlight
                cell.font = font_data_bold
            elif is_even:
                cell.fill = fill_zebra

    # ==========================================
    # HOJA 2: Supervisores vs Coordinadores de Proyecto
    # ==========================================
    ws2 = wb.create_sheet(title="2. Sup_vs_Coord_Contratista")
    ws2.views.sheetView[0].showGridLines = True
    
    headers_2 = [
        "N°", "ID SUPERVISOR", "SUPERVISOR TECSUR", "CR", 
        "EMPRESA SUBCONTRATISTA (12 OFICIALES)", 
        "DNI COORDINADOR (8 DÍGITOS OBLIGATORIO)", 
        "NOMBRE COMPLETO DEL COORDINADOR DE PROYECTO", 
        "TELÉFONO / CELULAR", "CORREO ELECTRÓNICO CONTRATISTA", "OBSERVACIONES / FRENTE"
    ]
    
    ws2.append(headers_2)
    ws2.row_dimensions[1].height = 28
    for col_num in range(1, len(headers_2) + 1):
        cell = ws2.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = align_center

    for idx, (sid, nombre, cargo, cr, crdesc) in enumerate(supervisores, start=1):
        contratista_sug = contratistas[(idx - 1) % len(contratistas)]
        row_vals = [
            idx, sid, nombre, cr, contratista_sug, "", "", "", "", "Frente habitual"
        ]
        ws2.append(row_vals)
        r = idx + 1
        ws2.row_dimensions[r].height = 20
        is_even = (idx % 2 == 0)
        
        for c in range(1, len(headers_2) + 1):
            cell = ws2.cell(row=r, column=c)
            cell.font = font_data
            cell.border = thin_border
            if c in (1, 2, 4):
                cell.alignment = align_center
            else:
                cell.alignment = align_left
            
            # Columnas a llenar por Edder / Jefaturas
            if c in (6, 7, 8, 9):
                cell.fill = fill_highlight
            elif is_even:
                cell.fill = fill_zebra

    # ==========================================
    # HOJA 3: Supervisores vs Capataces de Campo
    # ==========================================
    ws3 = wb.create_sheet(title="3. Sup_vs_Capataces_Campo")
    ws3.views.sheetView[0].showGridLines = True
    
    headers_3 = [
        "N°", "ID SUPERVISOR", "SUPERVISOR TECSUR", "CR", 
        "EMPRESA SUBCONTRATISTA", 
        "DNI CAPATAZ (8 DÍGITOS OBLIGATORIO)", 
        "NOMBRE COMPLETO DEL CAPATAZ", 
        "TIPO DE CUADRILLA HABITUAL", 
        "CELULAR CAPATAZ", "ZONA / FRENTE DE OPERACIÓN"
    ]
    
    ws3.append(headers_3)
    ws3.row_dimensions[1].height = 28
    for col_num in range(1, len(headers_3) + 1):
        cell = ws3.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = align_center

    for idx, (sid, nombre, cargo, cr, crdesc) in enumerate(supervisores, start=1):
        contratista_sug = contratistas[(idx - 1) % len(contratistas)]
        tipo_cuadrilla = "Planificación de Obras" if cr == '4070' else "Redes BT / Empalmes y Derivaciones"
        row_vals = [
            idx, sid, nombre, cr, contratista_sug, "", "", tipo_cuadrilla, "", "Sector asignado"
        ]
        ws3.append(row_vals)
        r = idx + 1
        ws3.row_dimensions[r].height = 20
        is_even = (idx % 2 == 0)
        
        for c in range(1, len(headers_3) + 1):
            cell = ws3.cell(row=r, column=c)
            cell.font = font_data
            cell.border = thin_border
            if c in (1, 2, 4):
                cell.alignment = align_center
            else:
                cell.alignment = align_left
            
            if c in (6, 7, 9):
                cell.fill = fill_highlight
            elif is_even:
                cell.fill = fill_zebra

    # ==========================================
    # HOJA 4: Supervisores vs Luz del Sur (LDS)
    # ==========================================
    ws4 = wb.create_sheet(title="4. Sup_vs_LuzDelSur")
    ws4.views.sheetView[0].showGridLines = True
    
    headers_4 = [
        "N°", "ID SUPERVISOR", "SUPERVISOR TECSUR", "CR", 
        "SECTOR / DISTRITO LUZ DEL SUR", 
        "INSPECTOR LUZ DEL SUR ASIGNADO", 
        "PROYECTISTA LDS RESPONSABLE", 
        "CORREO / TELÉFONO LDS", "TIPO DE OBRA PREDOMINANTE"
    ]
    
    ws4.append(headers_4)
    ws4.row_dimensions[1].height = 28
    for col_num in range(1, len(headers_4) + 1):
        cell = ws4.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = align_center

    for idx, (sid, nombre, cargo, cr, crdesc) in enumerate(supervisores, start=1):
        sector_sug = "Sede Central LDS" if cr == '4070' else "Sector Lima Sur / Chorrillos - SJM - Surco"
        tipo_obra = "Estudios y Proyectos BT" if cr == '4070' else "Obras Nuevas / Remodelación Redes BT"
        row_vals = [
            idx, sid, nombre, cr, sector_sug, "", "", "", tipo_obra
        ]
        ws4.append(row_vals)
        r = idx + 1
        ws4.row_dimensions[r].height = 20
        is_even = (idx % 2 == 0)
        
        for c in range(1, len(headers_4) + 1):
            cell = ws4.cell(row=r, column=c)
            cell.font = font_data
            cell.border = thin_border
            if c in (1, 2, 4):
                cell.alignment = align_center
            else:
                cell.alignment = align_left
            
            if c in (6, 7, 8):
                cell.fill = fill_highlight
            elif is_even:
                cell.fill = fill_zebra

    # ==========================================
    # HOJA 5: Catálogo 12 Contratistas Homologadas
    # ==========================================
    ws5 = wb.create_sheet(title="5. Directorio_12_Contratistas")
    ws5.views.sheetView[0].showGridLines = True
    
    headers_5 = [
        "N°", "RAZÓN SOCIAL DE LA EMPRESA", "RUC (OPCIONAL)", 
        "ESTADO DE HOMOLOGACIÓN", "ALCANCE OPERATIVO", 
        "REPRESENTANTE / CONTACTO PRINCIPAL", "TELÉFONO DE CONTACTO"
    ]
    
    ws5.append(headers_5)
    ws5.row_dimensions[1].height = 28
    for col_num in range(1, len(headers_5) + 1):
        cell = ws5.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = align_center

    for idx, contratista in enumerate(contratistas, start=1):
        row_vals = [
            idx, contratista, "", "Homologada Tecsur / LDS", "Obras de Baja Tensión y Redes", "", ""
        ]
        ws5.append(row_vals)
        r = idx + 1
        ws5.row_dimensions[r].height = 20
        is_even = (idx % 2 == 0)
        
        for c in range(1, len(headers_5) + 1):
            cell = ws5.cell(row=r, column=c)
            cell.font = font_data
            cell.border = thin_border
            if c in (1, 3, 4):
                cell.alignment = align_center
            else:
                cell.alignment = align_left
            if is_even:
                cell.fill = fill_zebra

    # Ajustar anchos automáticos en todas las hojas
    for ws in [ws0, ws1, ws2, ws3, ws4, ws5]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len and '\n' not in val:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Anchos específicos para Hoja 0
    ws0.column_dimensions['A'].width = 24
    ws0.column_dimensions['B'].width = 80

    # Guardar en la carpeta del proyecto y en el Escritorio
    path_proyecto = r"C:\Users\vllamocca\Desktop\PROGRAMA 20-20-20\REQUERIMIENTO_MATRIZ_ASIGNACIONES_EDDER_TECSUR.xlsx"
    path_escritorio = r"C:\Users\vllamocca\Desktop\REQUERIMIENTO_MATRIZ_ASIGNACIONES_EDDER_TECSUR.xlsx"
    
    wb.save(path_proyecto)
    wb.save(path_escritorio)
    
    print(f"Archivo guardado exitosamente en:")
    print(f"1) {path_proyecto}")
    print(f"2) {path_escritorio}")

if __name__ == "__main__":
    crear_excel_requerimiento_edder()
